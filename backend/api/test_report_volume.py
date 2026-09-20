from django.test import Client, TestCase

from api.models import (
    Application, Scholarship, StudentProfile, SystemSettings, User,
)

THIS_TERM = {'term_label': '26-1', 'school_year': '2026-2027',
             'semester': '1st Semester'}
LAST_TERM = {'term_label': '25-2', 'school_year': '2025-2026',
             'semester': '2nd Semester'}


def _student(n):
    user = User.objects.create_user(
        username=f'scholar{n:03d}@bipsu.edu.ph',
        email=f'scholar{n:03d}@bipsu.edu.ph', password='pw',
        first_name=f'Student{n:03d}', last_name=f'Surname{n:03d}',
        role='student', verification_status='approved')
    return StudentProfile.objects.create(
        user=user, student_id=f'2024-{n:05d}', course='BS Computer Science',
        year_level=(n % 4) + 1, gender='Female' if n % 2 else 'Male',
        municipality='Naval', province='Biliran')


class LengthyReportTest(TestCase):
    """The masterlist has to survive a realistic number of scholars.

    Asked for directly by the evaluators, who pointed out that the reports had
    only ever been exercised against a handful of rows. Sixty is past the point
    where a page break, a column width calculation or a query issued per row
    starts to matter, and still small enough to run on every build.
    """

    SCHOLARS = 60

    @classmethod
    def setUpTestData(cls):
        SystemSettings.objects.update_or_create(
            pk=1, defaults={'academic_year': '26-1',
                            'active_semester': '1st Semester'})
        cls.scholarship = Scholarship.objects.create(
            name='Academic Scholarship', type='Academic',
            category='application', description='x', eligibility='x',
            requirements=[])
        Application.objects.bulk_create([
            Application(student=_student(n), scholarship=cls.scholarship,
                        status='Approved', **THIS_TERM)
            for n in range(cls.SCHOLARS)
        ])
        cls.office = User.objects.create_user(
            username='office@bipsu.edu.ph', email='office@bipsu.edu.ph',
            password='pw', first_name='R', last_name='B', role='vpsea')

    def setUp(self):
        self.c = Client()
        self.assertTrue(self.c.login(email='office@bipsu.edu.ph', password='pw'))

    def test_the_reports_page_counts_every_scholar(self):
        """The page itself is a summary of counts; the names are in the files."""
        r = self.c.get('/vpsea/reports/?sy=26-1')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(
            r.context['grand_total'], self.SCHOLARS,
            f'the summary counted {r.context["grand_total"]} of '
            f'{self.SCHOLARS} scholars')

    def test_the_query_count_does_not_grow_with_the_scholar_list(self):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext

        with CaptureQueriesContext(connection) as small:
            self.c.get('/vpsea/reports/?sy=26-1')
        before = len(small.captured_queries)

        Application.objects.bulk_create([
            Application(student=_student(n), scholarship=self.scholarship,
                        status='Approved', **THIS_TERM)
            for n in range(self.SCHOLARS, self.SCHOLARS * 2)
        ])

        with CaptureQueriesContext(connection) as large:
            self.c.get('/vpsea/reports/?sy=26-1')
        after = len(large.captured_queries)

        self.assertLessEqual(
            after, before + 3,
            f'doubling the scholar list took the query count from {before} to '
            f'{after} — the page issues queries per row')

    def test_the_excel_download_carries_every_scholar(self):
        from io import BytesIO

        import openpyxl

        r = self.c.get('/vpsea/reports/download/excel/?sy=26-1')
        self.assertEqual(r.status_code, 200)
        book = openpyxl.load_workbook(BytesIO(r.content))
        found = {
            str(cell.value) for sheet in book.worksheets
            for row in sheet.iter_rows() for cell in row
            if cell.value is not None
        }
        text = ' '.join(found)
        for n in (0, self.SCHOLARS - 1):
            self.assertIn(f'Surname{n:03d}', text,
                          f'scholar {n} is missing from the spreadsheet')

    def test_the_pdf_preview_renders_at_this_size(self):
        r = self.c.get('/vpsea/reports/preview/?sy=26-1')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Type'], 'application/pdf')
        self.assertGreater(
            len(r.content), 2000,
            'the PDF came back too small to contain sixty scholars')


class ReportTermScopeTest(TestCase):
    """A report for one term must not contain another term's scholars.

    The imported rows were already filtered by term while the portal
    applications were not, so changing the term selector moved only part of
    the report. Anyone reading a semester's list was seeing every approved
    application ever recorded.
    """

    @classmethod
    def setUpTestData(cls):
        SystemSettings.objects.update_or_create(
            pk=1, defaults={'academic_year': '26-1',
                            'active_semester': '1st Semester'})
        scholarship = Scholarship.objects.create(
            name='Academic Scholarship', type='Academic',
            category='application', description='x', eligibility='x',
            requirements=[])
        cls.now = _student(1)
        cls.then = _student(2)
        Application.objects.create(student=cls.now, scholarship=scholarship,
                                   status='Approved', **THIS_TERM)
        Application.objects.create(student=cls.then, scholarship=scholarship,
                                   status='Approved', **LAST_TERM)
        User.objects.create_user(
            username='office@bipsu.edu.ph', email='office@bipsu.edu.ph',
            password='pw', first_name='R', last_name='B', role='vpsea')

    def setUp(self):
        self.c = Client()
        self.assertTrue(self.c.login(email='office@bipsu.edu.ph', password='pw'))

    def _names_in_spreadsheet(self, term):
        from io import BytesIO

        import openpyxl

        r = self.c.get(f'/vpsea/reports/download/excel/?sy={term}')
        self.assertEqual(r.status_code, 200)
        book = openpyxl.load_workbook(BytesIO(r.content))
        return ' '.join(
            str(cell.value) for sheet in book.worksheets
            for row in sheet.iter_rows() for cell in row
            if cell.value is not None)

    def test_both_terms_are_offered_in_the_selector(self):
        offered = dict(self.c.get('/vpsea/reports/').context['all_sy_display'])
        self.assertIn('26-1', offered)
        self.assertIn(
            '25-2', offered,
            'a term whose scholars came through the portal was not offered, '
            'so the office cannot report on it at all')

    def test_each_term_counts_only_its_own_scholar(self):
        for term in ('26-1', '25-2'):
            with self.subTest(term=term):
                total = self.c.get(f'/vpsea/reports/?sy={term}').context['grand_total']
                self.assertEqual(
                    total, 1,
                    f'{term} counted {total} scholars; each term has exactly one')

    def test_this_terms_spreadsheet_omits_last_terms_scholar(self):
        names = self._names_in_spreadsheet('26-1')
        self.assertIn('Surname001', names)
        self.assertNotIn(
            'Surname002', names,
            "last term's scholar appeared in this term's spreadsheet")

    def test_last_terms_spreadsheet_omits_this_terms_scholar(self):
        names = self._names_in_spreadsheet('25-2')
        self.assertIn('Surname002', names)
        self.assertNotIn(
            'Surname001', names,
            "this term's scholar appeared in last term's spreadsheet")
