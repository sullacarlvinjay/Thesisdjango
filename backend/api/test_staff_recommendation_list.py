from django.test import Client, TestCase

from api.models import (ApplicantRecord, Scholarship, StaffProfile,
                        SystemSettings, User)
from api.views_ranking import _staff_ranking_data


class RecommendationListHoldsOnlyTheQualifiedTest(TestCase):
    """The recommendation list is what the office sends up.

    A refusal and an application nobody can decide yet each belong on the
    page, with their reason, but not in the list headed "Recommendation" -
    the office reads that one as the answer.
    """

    def setUp(self):
        SystemSettings.objects.update_or_create(
            pk=1, defaults={'academic_year': '26-2',
                            'active_semester': '2nd Semester'})
        Scholarship.objects.create(
            name='Staff Scholarship', type='Staff', category='application',
            description='x', eligibility='x', requirements=[],
            accepting_applications=True)
        self.office = User.objects.create_user(
            username='sdso@bipsu.edu.ph', email='sdso@bipsu.edu.ph',
            password='pw', first_name='Office', last_name='Staff',
            role='vpsea')
        self.c = Client()
        self.assertTrue(self.c.login(email='sdso@bipsu.edu.ph', password='pw'))

    def _employee(self, name, number, employment):
        record = ApplicantRecord.objects.create(
            full_name=name, email=f'{number}@bipsu.edu.ph',
            qualified_for='Staff', status='Pending Validation', course='BSIT',
            year_level=1, school_year='2026-2027', semester='2nd Semester',
            term_label='26-2')
        record.is_nsu_staff = True
        record.student_id = number
        record.employment_status = employment
        record.save()
        return record

    def _dependent_who_graduated(self):
        employee = User.objects.create_user(
            username='rosa@bipsu.edu.ph', email='rosa@bipsu.edu.ph',
            password='pw', first_name='Rosa', last_name='Delgado',
            role='nsu_staff')
        profile = StaffProfile.objects.create(
            user=employee, employee_id='EMP-0042')
        profile.employment_status = 'Regular'
        profile.save()
        record = ApplicantRecord.objects.create(
            full_name='Mila Delgado', email='mila@bipsu.edu.ph',
            qualified_for='Staff', status='Pending Validation', course='BSCS',
            year_level=2, school_year='2026-2027', semester='2nd Semester',
            term_label='26-2')
        record.is_nsu_dependent = True
        record.student_id = '2024-00311'
        record.staff_name = 'Rosa Delgado'
        record.staff_employee_id = 'EMP-0042'
        record.relationship_to_staff = 'Daughter'
        record.has_baccalaureate = True
        record.save()
        return record

    def _names(self, bucket):
        return [e.applicant_name for e in _staff_ranking_data()[bucket]]

    def test_a_regular_employee_is_recommended(self):
        self._employee('Ana Cruz', 'EMP-0001', 'Regular')
        self.assertEqual(self._names('rows'), ['Ana Cruz'])

    def test_a_non_regular_employee_is_not_in_the_recommendation(self):
        self._employee('Ben Lim', 'EMP-0002', 'Job Order')
        self.assertNotIn('Ben Lim', self._names('rows'))
        self.assertIn('Ben Lim', self._names('refused'))

    def test_an_undecidable_application_is_in_neither(self):
        self._employee('Cely Uy', 'EMP-0003', '')
        self.assertNotIn('Cely Uy', self._names('rows'))
        self.assertNotIn('Cely Uy', self._names('refused'))
        self.assertIn('Cely Uy', self._names('needs_info'))

    def test_ranks_number_the_recommendation_from_one(self):
        self._employee('Ana Cruz', 'EMP-0001', 'Regular')
        self._employee('Dex Uy', 'EMP-0004', 'Regular')
        self._employee('Ben Lim', 'EMP-0002', 'Job Order')
        data = _staff_ranking_data()
        self.assertEqual([e.rank for e in data['rows']], [1, 2])
        self.assertEqual([e.rank for e in data['refused']], [None])

    def test_a_refusal_carries_the_rule_that_failed(self):
        self._employee('Ben Lim', 'EMP-0002', 'Job Order')
        refused = _staff_ranking_data()['refused'][0]
        self.assertEqual([r.key for r in refused.refusals], ['permanent'])
        self.assertIn('Job Order', refused.refusals[0].detail)

    def test_an_undecidable_one_reports_no_refusal(self):
        self._employee('Cely Uy', 'EMP-0003', '')
        undecided = _staff_ranking_data()['needs_info'][0]
        self.assertEqual(
            undecided.refusals, [],
            'a question nobody asked was reported as a reason for refusal')

    def test_a_graduated_dependent_is_refused_not_recommended(self):
        self._dependent_who_graduated()
        self.assertNotIn('Mila Delgado', self._names('rows'))
        refused = _staff_ranking_data()['refused'][0]
        self.assertEqual([r.key for r in refused.refusals], ['baccalaureate'])

    def test_the_counts_are_unchanged_by_the_split(self):
        self._employee('Ana Cruz', 'EMP-0001', 'Regular')
        self._employee('Ben Lim', 'EMP-0002', 'Job Order')
        self._employee('Cely Uy', 'EMP-0003', '')
        counts = _staff_ranking_data()['counts']
        self.assertEqual(counts['qualified'], 1)
        self.assertEqual(counts['not_qualified'], 1)
        self.assertEqual(counts['verification'], 1)
        self.assertEqual(_staff_ranking_data()['total'], 3)

    def test_the_page_shows_every_applicant_somewhere(self):
        self._employee('Ana Cruz', 'EMP-0001', 'Regular')
        self._employee('Ben Lim', 'EMP-0002', 'Job Order')
        self._employee('Cely Uy', 'EMP-0003', '')
        page = self.c.get('/vpsea/ranking/?type=Staff').content.decode()
        for name in ('Ana Cruz', 'Ben Lim', 'Cely Uy'):
            self.assertIn(name, page, f'{name} vanished from the page')
        self.assertIn('Applied, but not qualified', page)

    def test_the_download_carries_all_three(self):
        self._employee('Ana Cruz', 'EMP-0001', 'Regular')
        self._employee('Ben Lim', 'EMP-0002', 'Job Order')
        self._employee('Cely Uy', 'EMP-0003', '')
        response = self.c.get('/vpsea/ranking/download/?type=Staff')
        self.assertEqual(response.status_code, 200)

        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(response.content))
        text = '\n'.join(
            str(cell.value)
            for sheet in wb.worksheets
            for row in sheet.iter_rows()
            for cell in row if cell.value is not None)
        for name in ('Ana Cruz', 'Ben Lim', 'Cely Uy'):
            self.assertIn(
                name, text,
                f'{name} is on the page but missing from the download')


class StaffRankingStatusIsExhaustiveTest(TestCase):
    """Every verdict lands in exactly one of the three lists.

    The split is by status, so a fourth status added later would silently
    disappear from the page rather than fail anywhere. This is what notices.
    """

    def test_no_evaluation_falls_outside_the_three(self):
        SystemSettings.objects.update_or_create(
            pk=1, defaults={'academic_year': '26-2',
                            'active_semester': '2nd Semester'})
        for number, employment in (('EMP-1', 'Regular'), ('EMP-2', 'Part Time'),
                                   ('EMP-3', '')):
            record = ApplicantRecord.objects.create(
                full_name=f'Person {number}', email=f'{number}@bipsu.edu.ph',
                qualified_for='Staff', status='Pending Validation',
                course='BSIT', year_level=1, school_year='2026-2027',
                semester='2nd Semester', term_label='26-2')
            record.is_nsu_staff = True
            record.student_id = number
            record.employment_status = employment
            record.save()
        data = _staff_ranking_data()
        listed = len(data['rows']) + len(data['refused']) + len(data['needs_info'])
        self.assertEqual(
            listed, data['total'],
            'an application was screened but appears in none of the lists')
