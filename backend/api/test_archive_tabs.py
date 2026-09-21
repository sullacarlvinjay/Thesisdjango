import html as html_lib
import re

from django.test import Client, TestCase

from api.models import (
    Application, ImportedScholar, Scholarship, StudentProfile, SystemSettings,
    User,
)


def tab_hrefs(html):
    menu = re.search(r'id="archiveTabMenu".*?\n\s*</div>', html, re.S)
    found = re.findall(r'href="([^"]+)"', menu.group(0)) if menu else []
    return [html_lib.unescape(href) for href in found]


class ArchiveTabFixtures:
    term = '26-1'

    def setUp(self):
        SystemSettings.objects.update_or_create(
            pk=1, defaults={'academic_year': self.term,
                            'active_semester': '1st Semester'})
        self.ched = Scholarship.objects.create(
            name='CHED Merit', type='CHED', category='recommendation',
            group='external', description='x', eligibility='x', requirements=[])
        Scholarship.objects.create(
            name='Academic Scholarship', type='Academic', category='application',
            group='internal', description='x', eligibility='x', requirements=[])

        self.full = self.scholar('Reyes', '2022-00001', 'Full Merit / Full Scholar')
        self.half = self.scholar('Santos', '2022-00002', 'Half Merit / Partial Scholar')
        self.imported = ImportedScholar.objects.create(
            scholarship_type='CHED', term_label=self.term, last_name='Cruz',
            first_name='Juan', gender='M', course='BSIT', year_level=2)

        User.objects.create_user(
            username='v@bipsu.edu.ph', email='v@bipsu.edu.ph', password='pw',
            first_name='V', last_name='Officer', role='vpsea')
        self.c = Client()
        self.assertTrue(self.c.login(email='v@bipsu.edu.ph', password='pw'))

    def scholar(self, last, student_id, scholar_type):
        user = User.objects.create_user(
            username=f'{student_id}@bipsu.edu.ph', password='pw',
            email=f'{student_id}@bipsu.edu.ph', first_name='A', last_name=last,
            role='student')
        profile = StudentProfile.objects.create(
            user=user, student_id=student_id, course='BSCS', year_level=3,
            gender='Female', municipality='Naval')
        return Application.objects.create(
            student=profile, scholarship=self.ched, status='Approved',
            term_label=self.term, form_data={'scholar_type': scholar_type})

    def page(self, **query):
        response = self.c.get('/vpsea/archives/', query)
        self.assertEqual(response.status_code, 200)
        return response


class ChedIsTwoTabsTest(ArchiveTabFixtures, TestCase):
    def test_the_picker_offers_a_tab_for_each_ched_tier(self):
        hrefs = tab_hrefs(self.page(type='Academic').content.decode())
        self.assertIn('/vpsea/archives/?type=CHED&tier=Full', hrefs)
        self.assertIn('/vpsea/archives/?type=CHED&tier=Half', hrefs)

    def test_no_untiered_ched_tab_is_left_behind(self):
        hrefs = tab_hrefs(self.page(type='Academic').content.decode())
        self.assertNotIn('/vpsea/archives/?type=CHED', hrefs)

    def test_every_other_programme_keeps_exactly_one_tab(self):
        html = self.page(type='Academic').content.decode()
        hrefs = tab_hrefs(html)
        self.assertEqual(hrefs.count('/vpsea/archives/?type=Academic'), 1)

    def test_the_full_tab_lists_only_full_scholars(self):
        html = self.page(type='CHED', tier='Full').content.decode()
        self.assertIn('Reyes', html)
        self.assertNotIn('Santos', html)

    def test_the_half_tab_lists_only_half_scholars(self):
        html = self.page(type='CHED', tier='Half').content.decode()
        self.assertIn('Santos', html)
        self.assertNotIn('Reyes', html)

    def test_an_untiered_imported_row_still_prints_under_full(self):
        self.assertEqual(self.imported.award_tier, '')
        self.assertIn('Cruz', self.page(type='CHED', tier='Full').content.decode())

    def test_an_imported_row_stays_where_the_report_has_always_put_it(self):
        self.assertIn('Cruz', self.page(type='CHED', tier='Full').content.decode())
        self.assertNotIn('Cruz', self.page(type='CHED', tier='Half').content.decode())

    def test_an_untiered_award_prints_under_full_like_an_untiered_import(self):
        untiered = self.scholar('Bautista', '2022-00009', '')
        self.assertEqual(untiered.form_data['scholar_type'], '')

        self.assertIn('Bautista', self.page(type='CHED', tier='Full').content.decode())
        self.assertNotIn('Bautista', self.page(type='CHED', tier='Half').content.decode())

    def test_between_them_the_two_tabs_hold_every_ched_scholar(self):
        both = (self.page(type='CHED', tier='Full').content.decode()
                + self.page(type='CHED', tier='Half').content.decode())
        for name in ('Reyes', 'Santos', 'Cruz'):
            with self.subTest(name=name):
                self.assertIn(name, both)

    def test_a_bare_ched_url_lands_on_the_full_tab(self):
        response = self.page(type='CHED')
        self.assertEqual(response.context['active_tier'], 'Full')
        self.assertIn('Reyes', response.content.decode())

    def test_a_tier_that_is_not_a_tier_falls_back_to_full(self):
        self.assertEqual(self.page(type='CHED', tier='Platinum')
                         .context['active_tier'], 'Full')

    def test_a_tier_on_a_programme_that_has_none_is_ignored(self):
        response = self.page(type='Academic', tier='Full')
        self.assertEqual(response.context['active_tier'], '')
        self.assertEqual(response.context['active_tab_label'], 'Academic')

    def test_the_heading_names_the_tier_not_just_the_programme(self):
        self.assertEqual(self.page(type='CHED', tier='Half')
                         .context['active_tab_label'], 'CHED Half Merit')

    def test_the_block_band_is_dropped_once_the_tab_names_it(self):
        groups = self.page(type='CHED', tier='Full').context['scholar_groups']
        self.assertEqual(len(groups), 1)
        self.assertIsNone(groups[0]['title'])

    def listed_on(self, **kw):
        """The names the scholars table on one tab actually lists.

        Read off the table's own rows rather than searched for in the page.
        The page carries a scholar's name in more than one place — the recent
        activity card names whoever was added or deleted, whichever tab it
        happened on — so a page-wide search cannot tell "this tab lists her"
        from "something else on the page mentions her".
        """
        groups = self.page(**kw).context['scholar_groups']
        return ' '.join(row['search_name'] for group in groups
                        for row in group['rows']).lower()

    def test_a_scholar_added_on_the_full_tab_is_stored_as_full(self):
        response = self.c.post('/vpsea/archives/add/', {
            'scholarship_type': 'CHED', 'tier': 'Full',
            'first_name': 'Mia', 'last_name': 'Torres',
            'student_id': '2022-00003', 'course': 'BSCS', 'year_level': '1',
            'gender': 'Female', 'create_account': 'no',
        })
        self.assertEqual(response.status_code, 302)
        self.assertIn('tier=Full', response['Location'])
        self.assertEqual(
            ImportedScholar.objects.get(student_id='2022-00003').award_tier,
            'Full')
        self.assertIn('torres', self.listed_on(type='CHED', tier='Full'))
        self.assertNotIn('torres', self.listed_on(type='CHED', tier='Half'))

    def test_a_scholar_added_on_the_half_tab_is_stored_as_half(self):
        self.c.post('/vpsea/archives/add/', {
            'scholarship_type': 'CHED', 'tier': 'Half',
            'first_name': 'Mia', 'last_name': 'Torres',
            'student_id': '2022-00004', 'course': 'BSCS', 'year_level': '1',
            'gender': 'Female', 'create_account': 'no',
        })
        self.assertEqual(
            ImportedScholar.objects.get(student_id='2022-00004').award_tier,
            'Half')
        self.assertIn('torres', self.listed_on(type='CHED', tier='Half'))
        self.assertNotIn('torres', self.listed_on(type='CHED', tier='Full'))

    def test_the_download_link_on_a_tiered_tab_carries_the_tier(self):
        html = self.page(type='CHED', tier='Half').content.decode()
        self.assertIn('/vpsea/archives/download/?type=CHED&amp;tier=Half', html)

    def test_a_tiered_download_is_the_one_block(self):
        from io import BytesIO

        from openpyxl import load_workbook

        response = self.c.get('/vpsea/archives/download/',
                              {'type': 'CHED', 'tier': 'Full'})
        self.assertEqual(response.status_code, 200)
        rows = [list(r) for r in
                load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True)]
        values = {str(v) for row in rows for v in row if v is not None}
        self.assertIn('Reyes', values)
        self.assertNotIn('Santos', values)
        self.assertNotIn('Half Merit / Partial Scholar', values)


class ThePickerScalesTest(ArchiveTabFixtures, TestCase):
    def test_the_tab_strip_no_longer_wraps(self):
        html = self.page(type='Academic').content.decode()
        picker = html.split('<div class="tab-picker">')[1].split('<div class="flex gap-2">')[0]
        self.assertNotIn('flex-wrap', picker)
        self.assertIn('tab-picker__menu', picker)

    def test_every_archive_type_is_still_reachable(self):
        response = self.page(type='Academic')
        hrefs = ' '.join(tab_hrefs(response.content.decode()))
        for stype in response.context['archive_types']:
            with self.subTest(type=stype):
                if stype == 'CHED':
                    self.assertIn('type=CHED&tier=', hrefs)
                else:
                    self.assertIn(f'type={stype.replace(" ", "%20")}', hrefs)

    def test_a_newly_catalogued_programme_files_itself_under_its_funder(self):
        Scholarship.objects.create(
            name='Governor Award', type="Governor's Award", category='application',
            group='external', description='x', eligibility='x', requirements=[])
        groups = self.page(type='Academic').context['archive_tabs']
        external = next(g for g in groups if g['label'] == 'External')
        self.assertIn("Governor's Award", [t['label'] for t in external['tabs']])

    def test_the_groups_are_the_catalogues_own(self):
        groups = {g['label'] for g in self.page(type='Academic').context['archive_tabs']}
        self.assertIn('Internal', groups)
        self.assertIn('External', groups)

    def test_the_unawarded_tab_is_not_filed_as_a_programme(self):
        groups = self.page(type='Academic').context['archive_tabs']
        other = next(g for g in groups if g['label'] == 'Other')
        self.assertIn('No Scholarship', [t['label'] for t in other['tabs']])

    def test_exactly_one_tab_is_marked_current(self):
        for query in ({'type': 'Academic'}, {'type': 'CHED', 'tier': 'Half'}):
            with self.subTest(**query):
                groups = self.page(**query).context['archive_tabs']
                lit = [t for g in groups for t in g['tabs'] if t['active']]
                self.assertEqual(len(lit), 1)
