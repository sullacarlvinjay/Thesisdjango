from io import BytesIO

from openpyxl import Workbook

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase

from api.models import ImportedScholar, Scholarship, SystemSettings, User

CONTRACT = ['No.', 'Award Number', 'Last Name', 'First Name', 'Middle Name',
            'Sex', 'Brgy./St.', 'Municipality', 'Province', 'Congress District',
            'Course', 'Yr.', 'Scholarship Program']


def a_row(number, last, first):
    return [number, f'AW-{number}', last, first, 'R', 'F', 'Poblacion', 'Naval',
            'Biliran', '1st', 'BSCS', '2', 'DOST']


def workbook(sheets, headings=CONTRACT):
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(title=title)
        if headings is not None:
            ws.append(list(headings))
        for row in rows:
            ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(
        'dost.xlsx', buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


class ImportReadsEverySheetTest(TestCase):

    def setUp(self):
        SystemSettings.objects.update_or_create(
            pk=1, defaults={'academic_year': '26-1',
                            'active_semester': '1st Semester'})
        self.programme = Scholarship.objects.create(
            name='DOST Scholarship', type='DOST', category='application',
            description='x', eligibility='x', requirements=[])
        User.objects.create_user(
            username='v@bipsu.edu.ph', email='v@bipsu.edu.ph', password='pw',
            first_name='V', last_name='Officer', role='vpsea')
        self.c = Client()
        self.assertTrue(self.c.login(email='v@bipsu.edu.ph', password='pw'))

    def upload(self, file):
        return self.c.post('/vpsea/archives/import/', {
            'type': 'DOST', 'rollover_label': '26-1', 'file': file})

    def names(self):
        return sorted(ImportedScholar.objects.values_list('last_name', flat=True))

    def test_one_sheet_still_imports_exactly_as_before(self):
        self.upload(workbook([('Scholars', [a_row(1, 'Santos', 'Maria')])]))
        self.assertEqual(self.names(), ['Santos'])

    def test_a_second_sheet_is_no_longer_dropped(self):
        self.upload(workbook([
            ('1st Semester', [a_row(1, 'Santos', 'Maria')]),
            ('2nd Semester', [a_row(1, 'Cruz', 'Pedro')]),
        ]))
        self.assertEqual(self.names(), ['Cruz', 'Santos'])

    def test_every_sheet_of_a_thick_workbook_is_read(self):
        self.upload(workbook([
            ('BSCS', [a_row(1, 'Santos', 'Maria'), a_row(2, 'Cruz', 'Pedro')]),
            ('BSIT', [a_row(1, 'Reyes', 'Ana')]),
            ('BSED', [a_row(1, 'Lim', 'Jose'), a_row(2, 'Tan', 'Rosa')]),
        ]))
        self.assertEqual(self.names(), ['Cruz', 'Lim', 'Reyes', 'Santos', 'Tan'])
        self.assertEqual(ImportedScholar.objects.count(), 5)

    def test_the_count_the_office_is_told_covers_every_sheet(self):
        r = self.upload(workbook([
            ('A', [a_row(1, 'Santos', 'Maria')]),
            ('B', [a_row(1, 'Cruz', 'Pedro')]),
        ]))
        self.assertIn('import_ok=2', r['Location'])

    def test_a_sheet_that_holds_no_scholars_adds_nothing(self):
        self.upload(workbook([
            ('Scholars', [a_row(1, 'Santos', 'Maria')]),
            ('Notes', [['Prepared by the provincial office'], ['Totals below']]),
        ]))
        self.assertEqual(self.names(), ['Santos'])

    def test_an_empty_sheet_beside_a_full_one_is_harmless(self):
        self.upload(workbook([
            ('Scholars', [a_row(1, 'Santos', 'Maria')]),
            ('Blank', []),
        ]))
        self.assertEqual(self.names(), ['Santos'])

    def test_each_sheet_is_read_against_its_own_heading_row(self):
        self.programme.extra_columns = [
            {'key': 'extra_batch', 'label': 'Batch', 'type': 'text'}]
        self.programme.save(update_fields=['extra_columns'])

        wb = Workbook()
        wb.remove(wb.active)
        first = wb.create_sheet(title='With batch')
        first.append(CONTRACT + ['Batch'])
        first.append(a_row(1, 'Santos', 'Maria') + ['2026-A'])
        second = wb.create_sheet(title='Without batch')
        second.append(CONTRACT)
        second.append(a_row(1, 'Cruz', 'Pedro'))
        buf = BytesIO()
        wb.save(buf)
        self.upload(SimpleUploadedFile('dost.xlsx', buf.getvalue()))

        self.assertEqual(
            ImportedScholar.objects.get(last_name='Santos').extra_data,
            {'extra_batch': '2026-A'})
        self.assertEqual(
            ImportedScholar.objects.get(last_name='Cruz').extra_data, {})

    def test_a_sheet_the_workbook_left_hidden_is_still_read(self):
        wb = Workbook()
        wb.remove(wb.active)
        shown = wb.create_sheet(title='Shown')
        shown.append(CONTRACT)
        shown.append(a_row(1, 'Santos', 'Maria'))
        tucked = wb.create_sheet(title='Tucked away')
        tucked.append(CONTRACT)
        tucked.append(a_row(1, 'Cruz', 'Pedro'))
        tucked.sheet_state = 'hidden'
        buf = BytesIO()
        wb.save(buf)
        self.upload(SimpleUploadedFile('dost.xlsx', buf.getvalue()))
        self.assertEqual(self.names(), ['Cruz', 'Santos'])
