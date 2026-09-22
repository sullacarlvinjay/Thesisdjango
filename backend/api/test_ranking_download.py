from datetime import date
from io import BytesIO

from django.test import Client, TestCase
from openpyxl import load_workbook

from api.models import (
    AffirmativeRecommendation, ApplicantRecord, StudentProfile,
    SystemSettings, User,
)

URL = '/vpsea/ranking/download/'


class RankingDownloadFixtures:

    def setUp(self):
        SystemSettings.objects.create(pk=1, academic_year='26-1',
                                      active_semester='1st Semester')
        self.officer = User.objects.create_user(
            username='sdso@bipsu.edu.ph', email='sdso@bipsu.edu.ph',
            password='pw', role='vpsea', first_name='Ofelia', last_name='Reyes')
        self.c = Client()
        self.c.force_login(self.officer)

    def a_student(self, email='ana@bipsu.edu.ph', student_id='2022-00111',
                  first='Ana', last='Lim', **fields):
        user = User.objects.create_user(
            username=email, email=email, password='pw',
            first_name=first, last_name=last, role='student')
        return StudentProfile.objects.create(user=user, student_id=student_id, **fields)

    def an_eligible_student(self, **kw):
        return self.a_student(shs_gpa=91.0, suc_exam_score=42.0, suc_exam_total=50.0,
                              is_tes_beneficiary=False, course='BSCS', year_level=2,
                              **kw)

    TES_COMPLETE = dict(
        citizenship='Filipino',
        school='School of Technologies and Computer Studies',
        course='BSCS', year_level=2,
        has_previous_degree=False,
        family_income=120000.0, household_size=5,
        is_listahanan_household=False, is_4ps_beneficiary=False,
        is_solo_parent_dependent=False, disability_type='NO',
    )

    def a_ranked_student(self, **kw):
        fields = dict(self.TES_COMPLETE, year_first_enrolled=date.today().year - 1)
        fields.update(kw)
        return self.a_student(**fields)

    def book(self, tab='Affirmative', **params):
        query = '&'.join(f'{k}={v}' for k, v in params.items())
        response = self.c.get(f'{URL}?type={tab}' + (f'&{query}' if query else ''))
        self.assertEqual(response.status_code, 200)
        return response, load_workbook(BytesIO(response.content))

    def rows_of(self, ws):
        head = None
        for line in ws.iter_rows(values_only=True):
            if line[0] in ('Rank', 'Student', 'Applicant'):
                head = line
                continue
            if head is not None and line[0] is not None or (head and any(line)):
                yield dict(zip(head, line, strict=True))


class TheFileIsOfferedTest(RankingDownloadFixtures, TestCase):

    def test_each_tab_hands_back_a_workbook(self):
        for tab in ('Affirmative', 'TES', 'Staff'):
            response, _ = self.book(tab)
            self.assertIn('spreadsheetml', response['Content-Type'], tab)
            self.assertIn('attachment;', response['Content-Disposition'], tab)

    def test_the_file_is_named_for_its_tab_and_the_day(self):
        response, _ = self.book('TES')
        self.assertIn('BiPSU_TES_recommendation_', response['Content-Disposition'])
        self.assertIn('.xlsx', response['Content-Disposition'])
        self.assertNotIn('Recommendation_recommendation',
                         response['Content-Disposition'])

    def test_an_unknown_tab_falls_back_rather_than_failing(self):
        response, wb = self.book('Nonsense')
        self.assertEqual(wb.worksheets[0].title, 'Affirmative Action')

    def test_only_the_office_may_download_it(self):
        self.c.logout()
        self.assertEqual(self.c.get(f'{URL}?type=TES').status_code, 302)

    def test_a_student_may_not_download_the_list_they_are_on(self):
        student = self.an_eligible_student()
        self.c.force_login(student.user)
        self.assertEqual(self.c.get(f'{URL}?type=TES').status_code, 302)

    def test_the_page_offers_the_link_on_every_tab(self):
        for tab, expected in (('Affirmative', 'type=Affirmative'),
                              ('TES', 'type=TES'), ('Staff', 'type=Staff')):
            page = self.c.get(f'/vpsea/ranking/?type={tab}').content.decode()
            self.assertIn(f'/vpsea/ranking/download/?{expected}', page, tab)


class TheAffirmativeListTest(RankingDownloadFixtures, TestCase):

    def test_a_recommended_student_is_on_it_with_their_rank(self):
        self.an_eligible_student()
        self.c.get('/vpsea/ranking/')
        _, wb = self.book()
        rows = list(self.rows_of(wb.active))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['Student'], 'Ana Lim')
        self.assertEqual(rows[0]['Student No.'], '2022-00111')
        self.assertEqual(rows[0]['Rank'], 1)
        self.assertEqual(rows[0]['Status'], 'Recommended')

    def test_the_three_rules_are_three_columns_with_their_answers(self):
        self.an_eligible_student()
        self.c.get('/vpsea/ranking/')
        _, wb = self.book()
        row = next(self.rows_of(wb.active))
        self.assertEqual(row['GPA ≥ 75%'], 'Yes')
        self.assertEqual(row['Exam ≥ 50%'], 'Yes')
        self.assertEqual(row['Not a TES Beneficiary'], 'Yes')

    def test_the_threshold_on_the_page_is_the_threshold_in_the_file(self):
        self.an_eligible_student()
        self.c.get('/vpsea/ranking/')
        self.c.get('/vpsea/ranking/?passing=95')

        _, wb = self.book(passing=95)
        titles = [c.value for c in wb.active['A'][:6] if c.value]
        self.assertTrue(any('95%' in str(t) for t in titles), titles)
        row = next(self.rows_of(wb.active))
        self.assertEqual(row['GPA ≥ 95%'], 'No')
        self.assertIn(row['Rank'], (None, ''))

    def test_downloading_does_not_re_evaluate_anybody(self):
        self.an_eligible_student()
        self.c.get('/vpsea/ranking/')
        rec = AffirmativeRecommendation.objects.get()
        rec.status = 'Disqualified'
        rec.save(update_fields=['status'])

        self.book()
        rec.refresh_from_db()
        self.assertEqual(rec.status, 'Disqualified')


class TheTesListTest(RankingDownloadFixtures, TestCase):

    def test_every_ranked_student_is_on_it(self):
        self.a_ranked_student()
        self.a_ranked_student(email='b@bipsu.edu.ph', student_id='2022-00222',
                              first='Ben', last='Cruz')
        _, wb = self.book('TES')
        names = {row['Student'] for row in self.rows_of(wb.worksheets[0])}
        self.assertEqual(len(names), 2)

    def test_a_student_the_screen_held_back_is_a_count_not_a_row(self):
        self.a_ranked_student()
        self.a_student(email='c@bipsu.edu.ph', student_id='2022-00333',
                       first='Cita', last='Reyes')
        _, wb = self.book('TES')
        ws = wb.worksheets[0]
        names = {row['Student'] for row in self.rows_of(ws)}
        self.assertNotIn('Reyes, Cita', names)
        self.assertEqual(len(names), 1)
        titles = ' '.join(str(c.value) for c in ws['A'][:8] if c.value)
        self.assertIn('1 student(s) not shown', titles)

    def test_the_reasons_sheet_carries_the_why_panel(self):
        self.a_ranked_student()
        _, wb = self.book('TES')
        self.assertEqual(wb.worksheets[1].title, 'Reasons')
        rules = [row for row in self.rows_of(wb.worksheets[1])]
        labels = {row['Rule'] for row in rules}
        self.assertIn('Citizenship', labels)
        self.assertIn('Current College Enrollment', labels)
        for row in rules:
            self.assertTrue(row['Verdict'])
            self.assertTrue(row['Reading'])

    def test_no_status_is_written_onto_a_student(self):
        profile = self.an_eligible_student()
        self.book('TES')
        profile.refresh_from_db()
        self.assertFalse(profile.is_tes_beneficiary)


class TheStaffListTest(RankingDownloadFixtures, TestCase):

    def an_application(self, **fields):
        fields.setdefault('full_name', 'Earl Villablanca')
        fields.setdefault('email', 'earl@bipsu.edu.ph')
        fields.setdefault('qualified_for', 'Staff')
        fields.setdefault('status', 'Pending Validation')
        fields.setdefault('course', 'BSIT')
        return ApplicantRecord.objects.create(**fields)

    def test_a_qualified_applicant_is_on_it_with_their_verdict(self):
        self.an_application(is_nsu_staff=True, employment_status='Regular',
                            student_id='32-1-000111')
        _, wb = self.book('Staff')
        row = next(self.rows_of(wb.worksheets[0]))
        self.assertEqual(row['Applicant'], 'Earl Villablanca')
        self.assertEqual(row['Applying As'], 'Employee')
        self.assertEqual(row['Applied'], 'Yes')
        self.assertEqual(row['Eligibility'], 'Eligible')
        self.assertEqual(row['Rank'], 1)

    def test_an_undecidable_application_carries_no_rank(self):
        self.an_application(is_nsu_staff=True, employment_status='',
                            student_id='32-1-000222')
        _, wb = self.book('Staff')
        rows = [r for r in self.rows_of(wb.worksheets[0])
                if r['Eligibility'] == 'For Verification']
        self.assertTrue(rows, 'expected an application the rules cannot decide')
        for row in rows:
            self.assertIn(row['Rank'], (None, ''), row['Applicant'])
            self.assertTrue(row['Missing Information'])

    def test_the_reasons_sheet_names_the_qualification_and_its_source(self):
        self.an_application(is_nsu_staff=True, employment_status='Regular',
                            student_id='32-1-000111')
        _, wb = self.book('Staff')
        self.assertEqual(wb.worksheets[1].title, 'Reasons')
        rules = list(self.rows_of(wb.worksheets[1]))
        labels = {row['Qualification'] for row in rules}
        self.assertIn('Permanent appointment', labels)
        self.assertIn('No baccalaureate already', labels)
        for row in rules:
            self.assertTrue(row['Reading'])
