from io import BytesIO

from django.test import Client, TestCase
from openpyxl import load_workbook

from api.affirmative_ranking import target_groups
from api.models import AffirmativeRecommendation, StudentProfile, SystemSettings, User
from api.student_views import _affirmative_ranking_data
from api.fixtures_registration import a_student as a_registration


class Fixtures:
    def setUp(self):
        SystemSettings.objects.create(pk=1, academic_year='26-1',
                                      active_semester='1st Semester')
        self.officer = User.objects.create_user(
            username='sdso@bipsu.edu.ph', email='sdso@bipsu.edu.ph', password='pw',
            role='vpsea', first_name='Ofelia', last_name='Reyes')
        self.c = Client()
        self.c.force_login(self.officer)
        self.seq = 0

    def a_student(self, first='Ana', last='Lim', **fields):
        self.seq += 1
        email = f'{first}.{last}.{self.seq}@bipsu.edu.ph'.lower()
        user = User.objects.create_user(
            username=email, email=email, password='pw',
            first_name=first, last_name=last, role='student')
        fields.setdefault('disability_type', 'NO')
        fields.setdefault('highschool_is_public', False)
        fields.setdefault('is_from_depressed_area', False)
        return StudentProfile.objects.create(
            user=user, student_id=f'2022-{self.seq:05d}', **fields)

    def an_eligible_student(self, **fields):
        for name, value in (('shs_gpa', 91.0), ('suc_exam_score', 42.0),
                            ('suc_exam_total', 50.0), ('is_tes_beneficiary', False),
                            ('course', 'BSCS'), ('year_level', 2)):
            fields.setdefault(name, value)
        return self.a_student(**fields)

    def ranked(self):
        AffirmativeRecommendation.evaluate_and_sync(75.0)
        return _affirmative_ranking_data(75.0)

    def names_in_order(self):
        return [r['profile'].user.first_name for r in self.ranked()['rows']]


class TheFourGroupsTest(Fixtures, TestCase):

    def test_an_indigenous_group_is_one(self):
        groups = target_groups(self.a_student(indigenous_group='Manobo'))
        self.assertEqual(groups.count, 1)
        self.assertIn('Indigenous group (Manobo)', groups.markers)

    def test_a_disability_is_one(self):
        groups = target_groups(self.a_student(disability_type='Visual Disability'))
        self.assertIn('Person with a disability (Visual Disability)', groups.markers)

    def test_a_public_high_school_is_one(self):
        groups = target_groups(self.a_student(highschool_is_public=True))
        self.assertIn('Public school', groups.markers)

    def test_a_depressed_area_is_one(self):
        groups = target_groups(self.a_student(is_from_depressed_area=True))
        self.assertIn('Depressed area', groups.markers)

    def test_all_four_are_counted_together(self):
        groups = target_groups(self.a_student(
            indigenous_group='Manobo', disability_type='Visual Disability',
            highschool_is_public=True, is_from_depressed_area=True))
        self.assertEqual(groups.count, 4)

    def test_a_student_in_none_of_them_has_none(self):
        groups = target_groups(self.a_student())
        self.assertEqual(groups.count, 0)
        self.assertEqual(groups.summary, '')

    def test_declining_the_question_is_not_a_group(self):
        groups = target_groups(self.a_student(indigenous_group='N/A',
                                              disability_type='NO'))
        self.assertEqual(groups.count, 0)


class SilenceIsNotAnAnswerTest(Fixtures, TestCase):

    def test_an_unanswered_public_school_question_is_named_not_assumed(self):
        groups = target_groups(self.a_student(highschool_is_public=None))
        self.assertEqual(groups.count, 0)
        self.assertIn('Whether the high school attended was public', groups.unknown)

    def test_an_unanswered_depressed_area_question_is_named(self):
        groups = target_groups(self.a_student(is_from_depressed_area=None))
        self.assertIn('Whether the student is from a depressed area', groups.unknown)

    def test_a_record_taken_before_the_disability_question_is_named(self):
        groups = target_groups(self.a_student(disability_type=''))
        self.assertIn('Disability (or NO for none)', groups.unknown)

    def test_an_answered_no_is_not_chased(self):
        groups = target_groups(self.a_student(
            highschool_is_public=False, is_from_depressed_area=False,
            disability_type='NO'))
        self.assertEqual(groups.unknown, ())

    def test_a_blank_indigenous_group_is_not_chased(self):
        groups = target_groups(self.a_student(indigenous_group=''))
        self.assertEqual(groups.unknown, ())


class TheGroupsOrderTheListTest(Fixtures, TestCase):

    def test_a_target_group_outranks_a_higher_fit_score(self):
        self.an_eligible_student(first='Bea', shs_gpa=99.0)
        self.an_eligible_student(first='Ana', shs_gpa=76.0, highschool_is_public=True)
        self.assertEqual(self.names_in_order(), ['Ana', 'Bea'])

    def test_more_groups_outrank_fewer(self):
        self.an_eligible_student(first='Bea', shs_gpa=99.0, highschool_is_public=True)
        self.an_eligible_student(first='Ana', shs_gpa=76.0, highschool_is_public=True,
                                 is_from_depressed_area=True)
        self.assertEqual(self.names_in_order(), ['Ana', 'Bea'])

    def test_the_fit_score_still_separates_students_the_mandate_reaches_equally(self):
        self.an_eligible_student(first='Ana', shs_gpa=99.0, highschool_is_public=True)
        self.an_eligible_student(first='Bea', shs_gpa=76.0, highschool_is_public=True)
        self.assertEqual(self.names_in_order(), ['Ana', 'Bea'])

    def test_the_office_is_told_how_far_the_mandate_reaches(self):
        self.an_eligible_student(first='Ana', highschool_is_public=True)
        self.an_eligible_student(first='Bea')
        data = self.ranked()
        self.assertEqual(data['eligible_count'], 2)
        self.assertEqual(data['in_target_group_count'], 1)


class TheGroupsAreNotARuleTest(Fixtures, TestCase):

    def test_a_student_in_no_group_is_still_eligible(self):
        self.an_eligible_student(first='Bea')
        data = self.ranked()
        self.assertEqual(data['eligible_count'], 1)
        self.assertEqual(data['rows'][0]['rank'], 1)

    def test_every_group_does_not_rescue_a_failed_rule(self):
        student = self.an_eligible_student(
            first='Ana', indigenous_group='Manobo',
            disability_type='Visual Disability', highschool_is_public=True,
            is_from_depressed_area=True)
        self.assertEqual(self.ranked()['eligible_count'], 1)

        student.is_tes_beneficiary = True
        student.save()

        data = self.ranked()
        self.assertEqual(data['eligible_count'], 0)
        self.assertEqual(data['rows'][0]['groups'].count, 4)
        self.assertIsNone(data['rows'][0]['rank'])
        self.assertEqual(AffirmativeRecommendation.objects.get().status,
                         'Disqualified')

    def test_the_three_rules_are_still_the_only_ones_synced(self):
        self.an_eligible_student(first='Ana')
        created, _ = AffirmativeRecommendation.evaluate_and_sync(75.0)
        self.assertEqual(created, 1)
        self.assertEqual(AffirmativeRecommendation.objects.get().status,
                         'Recommended')


class ThePageShowsTheGroupsTest(Fixtures, TestCase):

    def page(self):
        return self.c.get('/vpsea/ranking/?type=Affirmative').content.decode()

    def test_the_column_carries_the_group_a_student_is_in(self):
        self.an_eligible_student(indigenous_group='Manobo')
        page = self.page()
        self.assertIn('Target Group', page)
        self.assertIn('Indigenous group (Manobo)', page)

    def test_the_page_says_what_it_is_ordered_on(self):
        self.an_eligible_student()
        page = self.page()
        self.assertIn('Grades are not the only', page)
        self.assertIn('tie-break', page)

    def test_an_unanswered_question_is_shown_as_unanswered(self):
        self.an_eligible_student(highschool_is_public=None,
                                 is_from_depressed_area=None)
        self.assertIn('2 unanswered', self.page())


class TheWorkbookCarriesTheGroupsTest(Fixtures, TestCase):

    def rows_of(self, ws):
        head = None
        for line in ws.iter_rows(values_only=True):
            if line[0] == 'Rank':
                head = line
                continue
            if head is not None and line[0] is not None:
                yield dict(zip(head, line, strict=True))

    def book(self):
        response = self.c.get('/vpsea/ranking/download/?type=Affirmative')
        self.assertEqual(response.status_code, 200)
        return load_workbook(BytesIO(response.content))

    def test_the_groups_are_columns_on_the_sheet(self):
        self.an_eligible_student(indigenous_group='Manobo', highschool_is_public=True)
        self.c.get('/vpsea/ranking/')
        row = next(self.rows_of(self.book().active))
        self.assertEqual(row['Groups Matched'], 2)
        self.assertIn('Indigenous group (Manobo)', row['Target Groups'])
        self.assertIn('Public school', row['Target Groups'])

    def test_an_unanswered_question_is_named_in_the_file_not_counted(self):
        self.an_eligible_student(is_from_depressed_area=None)
        self.c.get('/vpsea/ranking/')
        row = next(self.rows_of(self.book().active))
        self.assertIn('depressed area', row['Unanswered Group Questions'])

    def test_the_file_is_in_the_same_order_as_the_page(self):
        self.an_eligible_student(first='Bea', shs_gpa=99.0)
        self.an_eligible_student(first='Ana', shs_gpa=76.0, highschool_is_public=True)
        self.c.get('/vpsea/ranking/')
        filed = [r['Student'].split()[0] for r in self.rows_of(self.book().active)]
        self.assertEqual(filed, self.names_in_order())


class TheQuestionsAreAskedTest(TestCase):

    def test_a_registration_records_both_answers(self):
        Client().post('/register/', a_registration(
            email='ana@bipsu.edu.ph', student_id='23-0002',
            highschool_is_public='yes', is_from_depressed_area='yes'))
        profile = StudentProfile.objects.get(student_id='23-0002')
        self.assertIs(profile.highschool_is_public, True)
        self.assertIs(profile.is_from_depressed_area, True)

    def test_a_registration_that_skips_them_is_refused(self):
        payload = a_registration(email='ana@bipsu.edu.ph', student_id='23-0003')
        payload.pop('highschool_is_public')
        payload.pop('is_from_depressed_area')
        page = Client().post('/register/', payload).content.decode()
        self.assertIn('public school', page)
        self.assertIn('depressed area', page)
        self.assertFalse(StudentProfile.objects.filter(student_id='23-0003').exists())

    def test_no_is_recorded_as_no_and_not_as_unanswered(self):
        Client().post('/register/', a_registration(
            email='ana@bipsu.edu.ph', student_id='23-0004',
            highschool_is_public='no', is_from_depressed_area='no'))
        profile = StudentProfile.objects.get(student_id='23-0004')
        self.assertIs(profile.highschool_is_public, False)
        self.assertIs(profile.is_from_depressed_area, False)

    def test_the_registration_form_asks_both(self):
        page = Client().get('/register/').content.decode()
        self.assertIn('name="highschool_is_public"', page)
        self.assertIn('name="is_from_depressed_area"', page)


class AnExistingStudentCanStillAnswerTest(Fixtures, TestCase):
    def test_my_profile_takes_the_answer_after_the_school_names_are_locked(self):
        student = self.an_eligible_student(
            elementary='Naval Central School', highschool='Biliran NHS',
            last_school='Biliran NHS', highschool_is_public=None,
            is_from_depressed_area=None)
        self.c.force_login(student.user)
        self.c.post('/student/profile/', {
            'highschool_is_public': 'yes', 'is_from_depressed_area': 'yes',
            'disability_type': 'NO',
        })
        student.refresh_from_db()
        self.assertIs(student.highschool_is_public, True)
        self.assertIs(student.is_from_depressed_area, True)

    def test_the_locked_school_name_is_still_locked(self):
        student = self.an_eligible_student(
            elementary='Naval Central School', highschool='Biliran NHS',
            last_school='Biliran NHS')
        self.c.force_login(student.user)
        self.c.post('/student/profile/', {
            'highschool': 'Somewhere Else', 'disability_type': 'NO'})
        student.refresh_from_db()
        self.assertEqual(student.highschool, 'Biliran NHS')
