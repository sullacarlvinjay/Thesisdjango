import re
from datetime import date, datetime
from io import BytesIO

from django.test import Client, TestCase
from openpyxl import load_workbook

from api.models import (
    AffirmativeStaffApplication, Application, ImportedScholar, Scholarship,
    ScholarListImport, StudentProfile, SystemSettings, User,
)

URL = '/vpsea/archives/download/'


def headings_of(html, index=0):
    tables = re.findall(r'<table class="scholar-table".*?</thead>', html, re.S)
    if index >= len(tables):
        return []
    return [h.strip()
            for h in re.findall(r'<th(?:\s[^>]*)?>(.*?)</th>', tables[index], re.S)]


def rows_of(response_or_sheet):
    sheet = (response_or_sheet if hasattr(response_or_sheet, 'iter_rows')
             else load_workbook(BytesIO(response_or_sheet.content)).active)
    return [list(row) for row in sheet.iter_rows(values_only=True)
            if any(value not in (None, '') for value in row)]


def header_rows(rows):
    return [[value for value in row if value is not None]
            for row in rows if row and row[0] == 'No.']


def flat(rows):
    return {str(value) for row in rows for value in row if value is not None}


class ArchiveFixtures:
    term = '26-1'

    def setUp(self):
        SystemSettings.objects.update_or_create(
            pk=1, defaults={'academic_year': self.term,
                            'active_semester': '1st Semester'})
        self.programme = Scholarship.objects.create(
            name='Academic Scholarship', type='Academic', category='application',
            description='x', eligibility='x', requirements=[])

        user = User.objects.create_user(
            username='ana@bipsu.edu.ph', email='ana@bipsu.edu.ph', password='pw',
            first_name='Ana', last_name='Lim', role='student')
        self.profile = StudentProfile.objects.create(
            user=user, student_id='2022-00111', course='BSCS', year_level=3,
            gwa=1.25, gender='Female', municipality='Naval')
        self.award = Application.objects.create(
            student=self.profile, scholarship=self.programme, status='Approved',
            award_number='ACA-001', term_label=self.term)
        self.imported = ImportedScholar.objects.create(
            scholarship_type='Academic', term_label=self.term, last_name='Cruz',
            first_name='Juan', gender='M', course='BSIT', year_level=2, gwa=1.7,
            student_id='2021-00099')

        User.objects.create_user(
            username='v@bipsu.edu.ph', email='v@bipsu.edu.ph', password='pw',
            first_name='V', last_name='Officer', role='vpsea')
        self.c = Client()
        self.assertTrue(self.c.login(email='v@bipsu.edu.ph', password='pw'))

    def a_past_term(self, label, stype='Academic'):
        return ScholarListImport.objects.create(
            scholarship_type=stype, term_label=label, scholar_count=1,
            excel_file='rollovers/old.xlsx')

    def page(self, stype='Academic', **query):
        return self.c.get('/vpsea/archives/',
                          {'type': stype, **query}).content.decode()

    def sheet(self, stype='Academic', **query):
        response = self.c.get(URL, {'type': stype, **query})
        self.assertEqual(response.status_code, 200)
        return rows_of(response)


class TheFileSaysWhatTheTableSaysTest(ArchiveFixtures, TestCase):

    def test_the_headings_are_the_tables_own(self):
        on_screen = [h for h in headings_of(self.page()) if h != 'Actions']
        on_screen[0] = 'No.'
        self.assertEqual(header_rows(self.sheet())[0], on_screen)

    def test_rearranging_the_columns_rearranges_the_file(self):
        self.programme.table_columns = ['award_number', 'last_name', 'course']
        self.programme.save(update_fields=['table_columns'])
        self.assertEqual(header_rows(self.sheet())[0],
                         ['No.', 'Award No.', 'Last Name', 'Course'])

        self.programme.table_columns = ['course', 'award_number', 'last_name']
        self.programme.save(update_fields=['table_columns'])
        self.assertEqual(header_rows(self.sheet())[0],
                         ['No.', 'Course', 'Award No.', 'Last Name'])

    def test_an_imported_scholar_is_in_the_file(self):
        values = flat(self.sheet())
        self.assertIn('Cruz', values)
        self.assertIn('Lim', values)

    def test_the_rows_are_numbered_the_way_the_table_numbers_them(self):
        rows = self.sheet()
        self.assertEqual([row[0] for row in rows if isinstance(row[0], int)],
                         [1, 2])

    def test_the_file_follows_the_term_dropdown(self):
        self.a_past_term('25-2')
        ImportedScholar.objects.create(
            scholarship_type='Academic', term_label='25-2', last_name='Older',
            first_name='Term', gender='F', course='BSED', year_level=4)

        now = flat(self.sheet())
        self.assertIn('Cruz', now)
        self.assertNotIn('Older', now)

        then = flat(self.sheet(sy='25-2'))
        self.assertIn('Older', then)
        self.assertNotIn('Cruz', then)
        self.assertIn('Older', self.page(sy='25-2'), 'and the page agrees')

    def test_a_term_that_is_not_this_programmes_falls_back_like_the_page_does(self):
        self.assertIn('Cruz', flat(self.sheet(sy='not-a-term')))
        self.assertIn('Cruz', self.page(sy='not-a-term'))

    def test_ched_comes_down_in_the_two_blocks_the_page_shows(self):
        Scholarship.objects.create(name='CHED Merit', type='CHED',
                                   category='application', description='x',
                                   eligibility='x', requirements=[])
        rows = self.sheet('CHED')
        values = flat(rows)
        self.assertIn('Full Merit / Full Scholar', values)
        self.assertIn('Half Merit / Partial Scholar', values)
        headers = header_rows(rows)
        self.assertEqual(len(headers), 2, 'each block carries its own headings')
        self.assertEqual(headers[0], headers[1])

    def test_a_programme_the_old_download_had_no_branch_for_gets_its_own_columns(self):
        Scholarship.objects.create(
            name='Sports Scholarship', type='Sports', category='application',
            description='x', eligibility='x', requirements=[],
            table_columns=['last_name', 'course'])
        ImportedScholar.objects.create(
            scholarship_type='Sports', term_label=self.term, last_name='Dela Cruz',
            first_name='Pia', gender='F', course='BSHM', year_level=1)
        rows = self.sheet('Sports')
        self.assertEqual(header_rows(rows)[0], ['No.', 'Last Name', 'Course'])
        self.assertIn('Dela Cruz', flat(rows))

    def test_an_affirmative_or_staff_tab_downloads_its_own_records(self):
        for stype in ('Affirmative', 'Staff'):
            Scholarship.objects.create(name=f'{stype} Scholarship', type=stype,
                                       category='application', description='x',
                                       eligibility='x', requirements=[])
            AffirmativeStaffApplication.objects.create(
                full_name='Rosa Mendoza', contact_number='0918',
                date_of_birth='1990-01-01', course='BSED', year_level=1,
                status='Approved', qualified_for=stype, student_id='EMP-1')
            rows = self.sheet(stype)
            self.assertTrue(header_rows(rows), f'{stype} sheet had no headings')
            self.assertIn('Mendoza', flat(rows))

    def test_the_no_scholarship_tab_has_no_workbook_to_give(self):
        response = self.c.get(URL, {'type': 'No Scholarship'})
        self.assertEqual(response.status_code, 302)
        self.assertIn('/vpsea/archives/', response['Location'])

    def test_only_the_office_may_download(self):
        self.c.logout()
        self.assertTrue(self.c.login(email='ana@bipsu.edu.ph', password='pw'))
        self.assertNotEqual(self.c.get(URL, {'type': 'Academic'}).status_code, 200)


class TheOfficesOwnColumnsComeDownTooTest(ArchiveFixtures, TestCase):
    def setUp(self):
        super().setUp()
        self.programme.table_columns = ['last_name']
        self.programme.extra_columns = [
            {'key': 'extra_batch', 'label': 'Batch', 'type': 'text'},
            {'key': 'extra_stipend', 'label': 'Stipend', 'type': 'number'},
            {'key': 'extra_awarded_on', 'label': 'Awarded On', 'type': 'date'},
        ]
        self.programme.save(update_fields=['table_columns', 'extra_columns'])

    def test_a_custom_column_is_a_heading_in_the_file(self):
        self.assertEqual(header_rows(self.sheet())[0],
                         ['No.', 'Last Name', 'Batch', 'Stipend', 'Awarded On'])

    def test_the_values_typed_into_it_are_in_the_file(self):
        self.c.post('/vpsea/archives/columns/', {
            'type': 'Academic', 'sy': self.term,
            f'extra__award__{self.award.pk}__extra_batch': '2026-A',
            f'extra__imported__{self.imported.pk}__extra_batch': '2026-B',
        })
        values = flat(self.sheet())
        self.assertIn('2026-A', values)
        self.assertIn('2026-B', values)

    def test_a_number_arrives_as_a_number_and_a_date_as_a_date(self):
        self.c.post('/vpsea/archives/columns/', {
            'type': 'Academic', 'sy': self.term,
            f'extra__award__{self.award.pk}__extra_stipend': '2500',
            f'extra__award__{self.award.pk}__extra_awarded_on': '2026-06-15',
        })
        values = [value for row in self.sheet() for value in row]
        self.assertIn(2500, values, 'a Number column came down as text')
        dates = [value for value in values if isinstance(value, (date, datetime))]
        self.assertEqual([getattr(value, 'date', lambda: value)() for value in dates],
                         [date(2026, 6, 15)])
